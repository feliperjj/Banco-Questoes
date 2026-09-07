from pathlib import Path
import docx
import pytest
from src.importador.parser import parsear_questoes
from src.importador.extrator import extrair_texto_docx, extrair_texto_pdf
from src.importador.perfis.gabaritos_texto import ExtratorPadrao, ExtratorMultiprova
from src.importador.perfis.gabaritos_resultado import MapaGabarito
from src.importador.perfis.pdf_paginas import TextoExtraido, colunas_por_separadores


@pytest.mark.parametrize('marcador', ['QUESTÃO {n}\n', 'QUESTAO {n}\n', 'Questão nº {n}: ', 'Q. {n} - ', '{n}. ', '{n}) ', '{n}: ', '{n} - ', '{n}\n', '{n} '])
def test_formatos_preservam_numeracao_enunciado_e_alternativas(marcador):
    texto = '\n'.join(marcador.format(n=n)+'Escolha a definição correta para este problema.\nA) Primeira definição\nB) Segunda definição\nC) Terceira definição' for n in range(1,4))
    qs=parsear_questoes(texto)
    assert [q['numero'] for q in qs] == [1,2,3]
    assert all(q['enunciado']=='Escolha a definição correta para este problema.' for q in qs)
    assert all([a['letra'] for a in q['alternativas']]==list('ABC') for q in qs)


@pytest.mark.parametrize('opcoes', ['Ⓐ Uma opção\nⒷ Outra opção\nⒸ Terceira opção', 'A\nUma opção\nB\nOutra opção\nC\nTerceira opção', 'a) Uma opção b) Outra opção c) Terceira opção', '(A) Uma opção\n(B) Outra opção\n(C) Terceira opção'])
def test_variacoes_alternativas(opcoes):
    qs=parsear_questoes('QUESTÃO 1\nQual a resposta correta para este problema?\n'+opcoes)
    assert len(qs)==1
    assert [a['texto'] for a in qs[0]['alternativas']]==['Uma opção','Outra opção','Terceira opção']


def test_prova_mista_nao_descarta_item_de_julgamento():
    qs=parsear_questoes('QUESTÃO 1\nEscolha a resposta correta.\nA) Uma\nB) Outra\nQUESTÃO 2\nJulgue a afirmação: todos os arquivos são imutáveis.')
    assert [q['tipo'] for q in qs]==['multipla_escolha','certo_errado']


def test_intervalo_explicito_reutiliza_texto_sem_vazar_para_outra_questao():
    texto='Texto para as questões 1 a 2\n'+('Um texto comum com informações relevantes. '*5)+'\n'
    texto+='\n'.join(f'QUESTÃO {n}\nEscolha a afirmação correta.\nA) Uma\nB) Outra' for n in range(1,4))
    qs=parsear_questoes(texto)
    assert 'Um texto comum' in qs[0]['enunciado'] and 'Um texto comum' in qs[1]['enunciado']
    assert 'Um texto comum' not in qs[2]['enunciado']


def test_avisos_ocr_chegam_a_revisao():
    texto=TextoExtraido('QUESTÃO 1\nEscolha a afirmação correta.\nA) Uma\nB) Outra', 'pdf_ocr_v1', ['Confira a leitura OCR'])
    q=parsear_questoes(texto)[0]
    assert q['perfil_extracao']=='pdf_ocr_v1'
    assert 'Confira a leitura OCR' in q['aviso_importacao']


def test_docx_preserva_ordem_de_paragrafos_tabelas_e_celulas(tmp_path):
    d=docx.Document();d.add_paragraph('QUESTÃO 1');d.add_paragraph('Escolha a afirmação correta.')
    t=d.add_table(rows=2,cols=1);t.cell(0,0).text='A) Uma';t.cell(1,0).text='B) Outra'
    d.add_paragraph('QUESTÃO 2');d.add_paragraph('Julgue a afirmação: o texto foi preservado.')
    p=tmp_path/'prova.docx';d.save(p)
    qs=parsear_questoes(extrair_texto_docx(str(p)))
    assert [q['numero'] for q in qs]==[1,2]
    assert len(qs[0]['alternativas'])==2


@pytest.mark.parametrize('linhas,esperado', [
    (['1 Certo','2 Errado','3 Anulada'], {1:'Certo',2:'Errado',3:'Anulada'}),
    (['Item 1 2 3','Gabarito V F X'], {1:'Certo',2:'Errado',3:'Anulada'}),
    (['1','A','2','B'], {1:'A',2:'B'}),
    (['Item 1 2 3','Gabarito C X E'], {1:'C',2:'X',3:'E'}),
])
def test_perfis_de_gabarito(linhas,esperado):
    assert ExtratorPadrao('','').extrair(linhas)==esperado


def test_multiprova_sem_selecao_nao_escolhe_primeira():
    assert ExtratorMultiprova('').extrair(['1 A 2 B 1 C 2 D'])=={}


def test_conflito_na_fonte_nao_e_sobrescrito_nem_reinserido():
    mapa=MapaGabarito();mapa.incorporar({1:'A',2:'B'});mapa.incorporar({1:'C'});mapa.incorporar({1:'A'})
    assert mapa=={2:'B'} and mapa.conflitos=={1} and mapa.avisos


def test_tres_colunas_exigem_evidencia_geometrica():
    class Area:
        def __init__(self,n): self.n=n
        def extract_text(self,**kwargs):return f'{self.n}. Enunciado da questão\nA) Uma\nB) Outra'
    class Pagina:
        width=600; height=800
        lines=[dict(x0=x,x1=x,top=20,bottom=780) for x in (200,400)]
        rects=[]
        def extract_words(self,**kwargs):return []
        def crop(self,bbox):return Area(int(bbox[0]/200)+1)
    resultado=colunas_por_separadores(Pagina())
    assert resultado[1]=='pdf_3_colunas_separadas_v1'
    assert [q['numero'] for q in parsear_questoes(resultado[0])]==[1,2,3]


@pytest.mark.parametrize('nome', ['analista_area_ciencias_agrarias_subarea_sistemas_de_producao_animal.pdf','analista_area_ciencias_exatas_e_da_terra_subarea_sistemas_de_informacao.pdf'])
def test_embrapa_cem_itens_e_texto_base_entre_colunas(nome):
    p=Path(__file__).parents[1]/'samples/lote_2026_08_22'/nome
    if not p.exists(): pytest.skip('PDF local')
    qs=parsear_questoes(extrair_texto_pdf(str(p)),str(p))
    assert [q['numero'] for q in qs]==list(range(1,101))
    assert all(q['tipo']=='certo_errado' for q in qs)
    assert 'Presumivelmente' in qs[0]['enunciado']
    assert 'Presumivelmente' in qs[4]['enunciado']
    assert 'Climate change' in qs[12]['enunciado']
    assert 'Presumivelmente' not in qs[12]['enunciado']
    assert 'Climate change' not in qs[40]['enunciado']
    assert not any(q['enunciado'].endswith(('Conhecimentos Gerais','Conhecimentos Complementares')) for q in qs)


def test_conflito_na_mesma_pagina_nao_vira_resposta_da_ultima_prova():
    mapa=ExtratorPadrao('', '').extrair(['1 A 2 B', '1 C 2 B'])
    assert mapa == {2: 'B'}
    assert mapa.conflitos == {1}


def test_cargo_desconhecido_encerra_contexto_por_cabecalho_de_tipo():
    from src.importador.extrator import FiltroContexto
    f=FiltroContexto('PROVA TIPO 1','ADMINISTRADOR')
    assert f.pagina_relevante('ADMINISTRADOR – PROVA TIPO 1 - BRANCA\n1 A')
    assert not f.pagina_relevante('CONTADOR – PROVA TIPO 1 - BRANCA\n1 B')
    assert not f.pagina_relevante('1 C')
    assert not f.pagina_relevante('ADMINISTRADOR – PROVA TIPO 2 - VERDE\n1 D')


def test_anulacao_em_nota_nao_descarta_restante_do_gabarito():
    g=ExtratorPadrao('', '').extrair(['01: A 02: B 03: C', '*QUESTÃO 4: ANULADA'])
    assert g == {1:'A',2:'B',3:'C',4:'Anulada'}


def test_intervalo_no_titulo_nao_e_par_questao_resposta():
    g=ExtratorPadrao('', '').extrair(['Nível Superior - 1 a 28', '1 - E 2 - B'])
    assert g == {1:'E',2:'B'} and not g.conflitos


def test_titulo_de_cargo_quebrado_em_duas_linhas():
    from src.importador.extrator import FiltroContexto
    f=FiltroContexto(cargo='IBFC_04')
    linhas=['IBFC_04_VERSÃO A - ANALISTA DE SISTEMAS - PRODUÇÃO, REDES,',
            'SUPORTE DE BANCO DE DADOS E SUPORTE DE SISTEMAS', '1 2 3', 'A B C',
            'IBFC_05_OUTRO CARGO','1 2 3','B C D']
    assert f.linhas_do_cargo(linhas)==linhas[:4]


def test_cargos_distintos_na_mesma_pagina_preservam_o_selecionado():
    from src.importador.extrator import FiltroContexto
    f=FiltroContexto('PROVA TIPO 1','ADMINISTRADOR')
    texto='ADVOGADO – PROVA TIPO 1 - BRANCA\n1 A\nADMINISTRADOR – PROVA TIPO 1 - BRANCA\n1 B'
    assert f.pagina_relevante(texto)
    assert f.linhas_do_cargo(texto.splitlines())[-1]=='1 B'


def test_texto_rotulado_repetido_usa_ocorrencia_mais_proxima():
    texto='Texto 1\n'+('Um contexto sobre agricultura. '*5)+'\nQUESTÃO 1\nConsidere o texto 1 e assinale a resposta.\nA) Uma\nB) Outra\n'
    texto+='Texto 1\n'+('Outro contexto sobre computação. '*5)+'\nQUESTÃO 2\nConsidere o texto 1 e assinale a resposta.\nA) Uma\nB) Outra'
    qs=parsear_questoes(texto)
    assert 'Um contexto sobre agricultura' in qs[0]['enunciado']
    assert 'Outro contexto sobre computação' not in qs[0]['enunciado']
    assert 'Outro contexto sobre computação' in qs[1]['enunciado']


def test_grade_provas_acima_de_quatro_por_posicao_do_cabecalho():
    from src.importador.extrator import _extrair_grade_identificada
    class Pagina:
        def extract_words(self,**kwargs):
            return [dict(text=t,x0=x,x1=x+10,top=20) for x,t in [(10,'PROVA'),(24,'5'),(100,'PROVA'),(114,'6')]]
    texto='PROVA 5 PROVA 6\n21 - A 46 - B 21 - C 46 - D'
    assert _extrair_grade_identificada(Pagina(),texto,'PROVA 6','') == {21:'C',46:'D'}
    assert _extrair_grade_identificada(Pagina(),texto,'','') == {}


def test_caixas_em_duas_colunas_preservam_cinquenta_itens_fepese():
    p=Path(__file__).parents[1]/'samples/agente_de_tecnologia_da_informacao_e_comunicacao_analista_de_sistemas.pdf'
    if not p.exists(): pytest.skip('PDF local')
    qs=parsear_questoes(extrair_texto_pdf(str(p)),str(p))
    assert [q['numero'] for q in qs]==list(range(1,51))
    assert 'Taxa de desemprego' in qs[0]['enunciado']
    assert 'Extensão Rural é um processo' not in qs[0]['enunciado']
    assert all('SQUARE' not in a['texto'] for q in qs for a in q.get('alternativas') or [])


def test_docx_celula_mesclada_vertical_nao_duplica_conteudo(tmp_path):
    d=docx.Document();t=d.add_table(rows=2,cols=2)
    t.cell(0,0).merge(t.cell(1,0)).text='Texto compartilhado'
    t.cell(0,1).text='Primeiro';t.cell(1,1).text='Segundo'
    p=tmp_path/'mesclado.docx';d.save(p)
    texto=extrair_texto_docx(str(p))
    assert texto.splitlines()==['Texto compartilhado','Primeiro','Segundo']


def test_cargo_com_codigo_de_um_digito_encerra_bloco():
    from src.importador.extrator import FiltroContexto
    linhas=['2 - Auditor Fiscal','01 - A 02 - B','3 - Fonoaudiólogo','01 - C 02 - D']
    assert FiltroContexto(cargo='Auditor Fiscal').linhas_do_cargo(linhas)==linhas[:2]
